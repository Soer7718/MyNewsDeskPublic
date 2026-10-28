from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

def format_excel(filename):
    wb = load_workbook(filename)

    for ws in wb.worksheets:
        if ws.max_row < 2:
            continue

        # HEADER STYLE
        for cell in ws[1]:
            cell.font = Font(bold=True, size=12)
            cell.fill = PatternFill(start_color="D9D9D9", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # AUTO SIZE columns
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value else 0) for c in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 2

        # FREEZE PANES
        ws.freeze_panes = "A2"

        # THIN BORDERS
        thin = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = thin

        # SENTIMENT COLORS
        h = {str(ws.cell(1, c).value).lower(): c for c in range(1, ws.max_column + 1)}

        pos = h.get("sentiment_positive")
        neu = h.get("sentiment_neutral")
        neg = h.get("sentiment_negative")

        if pos:
            ws.conditional_formatting.add(
                f"{get_column_letter(pos)}2:{get_column_letter(pos)}{ws.max_row}",
                CellIsRule(operator="greaterThan", formula=["0"],
                           fill=PatternFill(start_color="C6EFCE", fill_type="solid"))
            )

        if neg:
            ws.conditional_formatting.add(
                f"{get_column_letter(neg)}2:{get_column_letter(neg)}{ws.max_row}",
                CellIsRule(operator="greaterThan", formula=["0"],
                           fill=PatternFill(start_color="F8CBAD", fill_type="solid"))
            )

        if neu:
            ws.conditional_formatting.add(
                f"{get_column_letter(neu)}2:{get_column_letter(neu)}{ws.max_row}",
                CellIsRule(operator="equal", formula=["0"],
                           fill=PatternFill(start_color="E7E6E6", fill_type="solid"))
            )

    wb.save(filename)
