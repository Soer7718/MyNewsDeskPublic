# src/pptx_updater.py
"""
Opdaterer 'Impact report - standard report sample.pptx' med et data-slide
der indeholder alle rapporter fra mynewsdesk_metrics.xlsx i en tabel.

Slidet tilføjes sidst i præsentationen og erstattes ved hver kørsel.

Brug:
    from src.pptx_updater import update_pptx
    update_pptx(excel_path, pptx_path)
"""

import re
from datetime import date
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

# Farver matchet til præsentationen
COLOR_DARK   = RGBColor(0x34, 0x30, 0x83)  # #343083 – mørk lilla (header/total)
COLOR_ACCENT = RGBColor(0x9B, 0x3C, 0xFA)  # #9B3CFA – lys lilla (accent)
COLOR_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_LIGHT  = RGBColor(0xF2, 0xF0, 0xFA)  # meget lys lilla (zebra)
COLOR_TEXT   = RGBColor(0x1A, 0x1A, 0x2E)  # næsten sort (brødtekst)
COLOR_GREY   = RGBColor(0x99, 0x99, 0x99)  # grå (footer)

# Markør bruges til at genkende og erstatte slidet ved næste kørsel
DATA_SLIDE_MARKER = "DATA_OVERVIEW_SLIDE_AUTO"

HEADERS = ["Rapport", "Coverage", "Impressions", "Ad Value (DKK)", "Pos.", "Neu.", "Neg."]


# ---------------------------------------------------------------------------
# Hjælpefunktioner
# ---------------------------------------------------------------------------

def _read_excel(excel_path: str) -> list:
    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active

    headers = [str(c.value).lower().strip() if c.value else "" for c in ws[1]]
    col = {h: i for i, h in enumerate(headers)}

    def _to_int(val):
        if val is None:
            return 0
        cleaned = re.sub(r"[^\d]", "", str(val))
        return int(cleaned) if cleaned else 0

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        rows.append([
            str(row[col.get("report_name", 0)] or ""),
            _to_int(row[col.get("coverage", 1)]),
            _to_int(row[col.get("article_impressions", 2)]),
            _to_int(row[col.get("ad_value_dkk", 3)]),
            _to_int(row[col.get("sentiment_positive", 4)]),
            _to_int(row[col.get("sentiment_neutral", 5)]),
            _to_int(row[col.get("sentiment_negative", 6)]),
        ])

    print(f"[pptx_updater] Læste {len(rows)} rapporter fra Excel")
    return rows


def _fmt(val: int) -> str:
    if val is None or val == 0:
        return "0"
    if val >= 1_000_000:
        return f"{val / 1_000_000:.1f}m".replace(".", ",")
    elif val >= 1_000:
        v = val / 1_000
        if v == int(v):
            return f"{int(v)}k"
        return f"{v:.1f}k".replace(".", ",")
    return str(val)


def _set_cell_text(cell, text, bold=False, font_size=9,
                   color=None, align=PP_ALIGN.LEFT):
    tf = cell.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.runs[0] if para.runs else para.add_run()
    run.text = str(text)
    run.font.bold = bold
    run.font.size = Pt(font_size)
    if color:
        run.font.color.rgb = color


def _set_cell_fill(cell, rgb: RGBColor):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for existing in tcPr.findall(qn('a:solidFill')):
        tcPr.remove(existing)
    solidFill = etree.SubElement(tcPr, qn('a:solidFill'))
    srgbClr = etree.SubElement(solidFill, qn('a:srgbClr'))
    srgbClr.set('val', f'{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}')


def _remove_existing_data_slide(prs: Presentation):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and DATA_SLIDE_MARKER in shape.text_frame.text:
                for sldId in prs.slides._sldIdLst:
                    r_id = sldId.get(
                        '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
                    )
                    if prs.slides.part.related_parts.get(r_id) is slide.part:
                        prs.slides._sldIdLst.remove(sldId)
                        print("[pptx_updater] Fjernede gammelt data-slide")
                        return


def _add_data_slide(prs: Presentation, rows: list):
    sw = prs.slide_width
    sh = prs.slide_height
    MARGIN = Cm(1.2)

    slide = prs.slides.add_slide(prs.slide_layouts[1])

    # Usynlig markør til genkendelse ved næste kørsel
    marker_box = slide.shapes.add_textbox(0, 0, Cm(0.1), Cm(0.1))
    marker_box.text_frame.paragraphs[0].add_run().text = DATA_SLIDE_MARKER
    marker_box.text_frame.paragraphs[0].runs[0].font.size = Pt(1)
    marker_box.text_frame.paragraphs[0].runs[0].font.color.rgb = COLOR_WHITE

    # TITEL
    title_box = slide.shapes.add_textbox(MARGIN, Cm(0.8), sw - 2 * MARGIN, Cm(1.4))
    para = title_box.text_frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = "Data Overview – All Reports"
    run.font.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = COLOR_DARK

    # DATO
    date_box = slide.shapes.add_textbox(MARGIN, Cm(2.0), sw - 2 * MARGIN, Cm(0.6))
    para2 = date_box.text_frame.paragraphs[0]
    para2.alignment = PP_ALIGN.CENTER
    run2 = para2.add_run()
    run2.text = f"Opdateret: {date.today().strftime('%d. %B %Y')}"
    run2.font.size = Pt(9)
    run2.font.color.rgb = COLOR_ACCENT
    run2.font.italic = True

    # TABEL
    n_rows = len(rows) + 2
    n_cols = len(HEADERS)
    tbl_top = Cm(2.8)
    tbl_h = sh - tbl_top - Cm(1.5)

    table = slide.shapes.add_table(
        n_rows, n_cols, MARGIN, tbl_top, sw - 2 * MARGIN, tbl_h
    ).table

    col_widths = [0.32, 0.10, 0.14, 0.15, 0.09, 0.09, 0.09]
    for i, w in enumerate(col_widths):
        table.columns[i].width = int((sw - 2 * MARGIN) * w)

    # Header
    for j, hdr in enumerate(HEADERS):
        cell = table.cell(0, j)
        _set_cell_fill(cell, COLOR_DARK)
        _set_cell_text(cell, hdr, bold=True, font_size=9, color=COLOR_WHITE,
                       align=PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT)

    # Data
    totals = [0] * 6
    for i, row in enumerate(rows):
        name = row[0]
        vals = row[1:]
        for k, v in enumerate(vals):
            totals[k] += v or 0
        bg = COLOR_LIGHT if i % 2 == 0 else COLOR_WHITE
        cells_data = [name] + [_fmt(v) for v in vals]
        for j, txt in enumerate(cells_data):
            cell = table.cell(i + 1, j)
            _set_cell_fill(cell, bg)
            _set_cell_text(cell, txt, bold=False, font_size=8, color=COLOR_TEXT,
                           align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT)

    # Total
    total_cells = ["TOTAL"] + [_fmt(v) for v in totals]
    for j, txt in enumerate(total_cells):
        cell = table.cell(n_rows - 1, j)
        _set_cell_fill(cell, COLOR_DARK)
        _set_cell_text(cell, txt, bold=True, font_size=9, color=COLOR_WHITE,
                       align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT)

    # Footer
    footer = slide.shapes.add_textbox(MARGIN, sh - Cm(1.2), sw - 2 * MARGIN, Cm(0.8))
    para3 = footer.text_frame.paragraphs[0]
    para3.alignment = PP_ALIGN.CENTER
    run3 = para3.add_run()
    run3.text = "STB x Related: Nordic Market Insights"
    run3.font.size = Pt(7)
    run3.font.color.rgb = COLOR_GREY

    print(f"[pptx_updater] Data-slide tilføjet med {len(rows)} rapporter")


# ---------------------------------------------------------------------------
# Hoved-funktion
# ---------------------------------------------------------------------------

def update_pptx(excel_path: str, pptx_path: str, output_path: str = None):
    """
    Læser Excel-filen og opdaterer PowerPoint med et nyt data-slide.

    Args:
        excel_path:  sti til mynewsdesk_metrics.xlsx
        pptx_path:   sti til den eksisterende PPTX
        output_path: hvor den opdaterede PPTX gemmes (default: overskriver pptx_path)
    """
    if output_path is None:
        output_path = pptx_path

    print(f"\n[pptx_updater] Læser Excel: {excel_path}")
    rows = _read_excel(excel_path)

    print(f"[pptx_updater] Åbner PPTX: {pptx_path}")
    prs = Presentation(pptx_path)

    _remove_existing_data_slide(prs)
    _add_data_slide(prs, rows)

    prs.save(output_path)
    print(f"[pptx_updater] PPTX gemt: {output_path}")
    return output_path
